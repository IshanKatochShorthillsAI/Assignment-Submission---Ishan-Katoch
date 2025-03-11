import unittest
import os
import sys
import sqlite3
import json
import csv
import tempfile
import re
import base64
from openpyxl import Workbook

# Add the src folder (and project root for detailed_test_case.py) to sys.path.
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from fileloader import PDFLoader, DOCXLoader, PPTLoader
from data_extractor import DataExtractor
from storage import FileStorage, SQLStorage
from detailed_test_case import DetailedTestCase


##########################################
# Dummy Classes to Support PPT Image Extraction
##########################################
class DummyImagePart:
    def __init__(self):
        self.blob = b"fakepptimagedata"
        self.ext = "jpeg"


class DummyPart:
    def __init__(self):
        self.related_parts = {"rIdTest": DummyImagePart()}


class FakeBlipElement:
    def get(self, key):
        # Always return "rIdTest" for the embed key.
        return "rIdTest"


class FakeElement:
    def xpath(self, query):
        # Return a list with one fake blip element.
        return [FakeBlipElement()]


##########################################
# Fake Classes for Testing
##########################################


# ----- Fake PDF Classes -----
class FakePDF:
    def __init__(self, metadata=None):
        self.page_count = 2
        self.metadata = (
            metadata
            if metadata is not None
            else {"title": "Fake PDF", "author": "Tester"}
        )

    def __getitem__(self, index):
        return FakePDFPage(index)


class FakePDFPage:
    def __init__(self, number):
        self._number = number
        self.rotation = 0
        self.rect = type("Rect", (), {"x0": 0, "y0": 0, "x1": 612, "y1": 792})

    @property
    def number(self):
        return self._number

    def get_text(self, mode):
        return {
            "blocks": [
                {
                    "type": 0,
                    "lines": [
                        {
                            "spans": [
                                {
                                    "text": f"Fake PDF text on page {self._number + 1}",
                                    "font": "FakeFont",
                                    "size": 12,
                                    "bbox": [0, 0, 100, 20],
                                    "flags": 0,
                                    "origin": [0, 0],
                                    "color": 0,
                                }
                            ]
                        }
                    ],
                }
            ]
        }

    def get_images(self, full=False):
        return []

    def extract_image(self, xref):
        return {"image": b"fakeimagedata", "width": 200, "height": 100, "ext": "png"}


class FakePDFLoader(PDFLoader):
    def __init__(self, metadata=None):
        self.fake_metadata = metadata

    def load_file(self, filepath: str):
        return FakePDF(metadata=self.fake_metadata)


# ----- Fake DOCX Classes -----
class FakeDOCX:
    def __init__(self, paragraphs=None, tables=None):
        self._paragraphs = (
            paragraphs if paragraphs is not None else [self.FakeParagraph()]
        )
        self.tables = tables if tables is not None else []

    @property
    def paragraphs(self):
        return self._paragraphs

    class FakeParagraph:
        def __init__(self, text="Fake DOCX text"):
            self.text = text
            self.style = "Normal"
            self.alignment = None
            self._runs = [self.FakeRun(text)]

        @property
        def runs(self):
            return self._runs

        class FakeRun:
            def __init__(self, text):
                self.text = text
                self.font = type(
                    "FakeFont",
                    (),
                    {
                        "name": "FakeFont",
                        "size": type("FakeSize", (), {"pt": 12})(),
                        "bold": False,
                        "italic": False,
                        "underline": False,
                        "color": type("FakeColor", (), {"rgb": "000000"})(),
                        "highlight_color": None,
                    },
                )()
                self.bold = False
                self.italic = False
                self.underline = False

    class FakeTable:
        def __init__(self, rows):
            self.style = "ComplexTableStyle"
            self.rows = [FakeDOCX.FakeTableRow(row) for row in rows]

    class FakeTableRow:
        def __init__(self, cells_text):
            self.cells = [FakeDOCX.FakeTableCell(text) for text in cells_text]

    class FakeTableCell:
        def __init__(self, text):
            self.text = text


class FakeDOCXLoader(DOCXLoader):
    def __init__(self, paragraphs=None, tables=None):
        self.fake_docx = FakeDOCX(paragraphs=paragraphs, tables=tables)

    def load_file(self, filepath: str):
        return self.fake_docx


# ----- Fake PPTX Classes -----
class FakePPT:
    def __init__(self, slides=None):
        self.slides = (
            slides if slides is not None else [FakePPTSlide(1), FakePPTSlide(2)]
        )


class FakePPTSlide:
    def __init__(self, slide_number):
        self.slide_number = slide_number
        self.shapes = []
        if slide_number == 1:
            self.shapes.append(
                FakePPTTextShape("Fake PPT text on slide 1", shape_id=101)
            )
            self.shapes.append(FakePPTTableShape(shape_id=102))
        else:
            self.shapes.append(
                FakePPTTextShape("Fake PPT text on slide 2", shape_id=201)
            )


class FakePPTTextShape:
    def __init__(self, text, shape_id):
        self._text = text
        self.shape_id = shape_id
        self.left = 100
        self.top = 100
        self.width = 500
        self.height = 200

    @property
    def text(self):
        return self._text

    @property
    def text_frame(self):
        return FakeTextFrame(self._text)


class FakeTextFrame:
    def __init__(self, text):
        self.paragraphs = [FakeTextParagraph(text)]


class FakeTextParagraph:
    def __init__(self, text):
        self.text = text
        self.alignment = 1
        self.runs = [FakeTextRun(text)]


class FakeTextRun:
    def __init__(self, text):
        self.text = text
        self.font = type(
            "FakeFont",
            (),
            {
                "name": "DefaultFont",
                "size": type("FakeSize", (), {"pt": 12})(),
                "bold": False,
                "italic": False,
                "underline": False,
                "color": type("FakeColor", (), {"rgb": "000000"})(),
                "highlight_color": None,
            },
        )()


class FakePPTTableShape:
    def __init__(self, shape_id):
        self.shape_id = shape_id
        self.left = 50
        self.top = 50
        self.width = 300
        self.height = 150
        self._has_table = True
        self.table = FakePPTTable()

    @property
    def has_table(self):
        return self._has_table


class FakePPTTable:
    def __init__(self):
        self.rows = [
            FakePPTTableRow(["Name", "Grade", "Remark"]),
            FakePPTTableRow(["Alice", "X", "Good"]),
        ]


class FakePPTTableRow:
    def __init__(self, cells_text):
        self.cells = [FakePPTTableCell(text) for text in cells_text]


class FakePPTTableCell:
    def __init__(self, text):
        self.text = text


class FakePPTLoader(PPTLoader):
    def load_file(self, filepath: str):
        return FakePPT()


# Extended Fake PPT Image Shape for testing image extraction.
class FakePPTImageShape:
    def __init__(self, shape_id):
        self.shape_id = shape_id
        self.left = 100
        self.top = 100
        self.width = 400
        self.height = 300
        self.part = DummyPart()

    @property
    def text(self):
        return ""

    @property
    def image(self):
        return type(
            "FakeImage",
            (),
            {
                "blob": b"fakepptimagedata",
                "ext": "jpeg",
                "width": 400,
                "height": 300,
            },
        )()

    @property
    def element(self):
        return FakeElement()


# Create a Fake PPT slide that includes an image shape.
class FakePPTSlideWithImage(FakePPTSlide):
    def __init__(self, slide_number):
        self.slide_number = slide_number
        self.shapes = [
            FakePPTTextShape("Slide with image", shape_id=301),
            FakePPTImageShape(shape_id=302),
        ]


# And a Fake PPT Loader that returns a presentation with an image.
class FakePPTLoaderWithImage(PPTLoader):
    def load_file(self, filepath: str):
        # Return an object with a slides attribute.
        return type("FakePPTWithImage", (), {"slides": [FakePPTSlideWithImage(1)]})()


# ----- Extended Fake for Error Cases -----
class InvalidFileLoader(PDFLoader):
    def load_file(self, filepath: str):
        raise ValueError("Invalid file type.")


class CorruptedPDFLoader(PDFLoader):
    def load_file(self, filepath: str):
        raise IOError("Corrupted file.")


##########################################
# Custom TestResult and Runner for Excel Report
##########################################
class ExcelTestResult(unittest.TextTestResult):
    def __init__(self, stream, descriptions, verbosity):
        super().__init__(stream, descriptions, verbosity)
        self.results = []  # List of (test_name, outcome, message)

    def addSuccess(self, test):
        super().addSuccess(test)
        detail = ""
        if hasattr(test, "get_detail"):
            detail = test.get_detail()
        self.results.append((self.getDescription(test), "PASS", detail))

    def addError(self, test, err):
        super().addError(test, err)
        detail = ""
        if hasattr(test, "get_detail"):
            detail = test.get_detail()
        detail += "\n" + self._exc_info_to_string(err, test)
        self.results.append((self.getDescription(test), "ERROR", detail))

    def addFailure(self, test, err):
        super().addFailure(test, err)
        detail = ""
        if hasattr(test, "get_detail"):
            detail = test.get_detail()
        detail += "\n" + self._exc_info_to_string(err, test)
        self.results.append((self.getDescription(test), "FAIL", detail))


class ExcelTestRunner(unittest.TextTestRunner):
    resultclass = ExcelTestResult

    def run(self, test):
        result = super().run(test)
        self.write_excel_report(result.results)
        return result

    def write_excel_report(self, results):
        wb = Workbook()
        ws = wb.active
        ws.title = "Test Results"
        ws.append(["Test Case", "Outcome", "Details"])
        for test_name, outcome, detail in results:
            ws.append([test_name, outcome, detail])
        report_file = os.path.join(os.getcwd(), "test_results.xlsx")
        wb.save(report_file)
        print(f"Excel test report saved to {report_file}")


##########################################
# Unit Tests
##########################################
class TestDataExtractor(DetailedTestCase):

    def test_pdf_extraction(self):
        extractor = DataExtractor(FakePDFLoader())
        data = extractor.extract_text("dummy.pdf")
        self.assertIn("document_metadata", data[0])
        page1 = next(item for item in data if item.get("page") == 1)
        self.record_detail(
            f"Expected: Fake PDF text on page 1, got: {page1.get('line_text')}"
        )
        self.assertEqual(page1.get("line_text"), "Fake PDF text on page 1")

    def test_pdf_missing_metadata(self):
        extractor = DataExtractor(FakePDFLoader(metadata={}))
        data = extractor.extract_text("dummy.pdf")
        self.assertIn("document_metadata", data[0])
        self.assertEqual(data[0]["document_metadata"], {})

    def test_docx_extraction(self):
        extractor = DataExtractor(FakeDOCXLoader())
        data = extractor.extract_text("dummy.docx")
        self.record_detail(f"Expected: Fake DOCX text, got: {data[0]['text']}")
        self.assertEqual(data[0]["text"], "Fake DOCX text")
        self.assertEqual(data[0]["runs"][0]["text"], "Fake DOCX text")

    def test_docx_complex_table(self):
        fake_table = FakeDOCX.FakeTable(
            [
                ["Header1", "Header2", "Header3"],
                ["Row1Col1", "Row1Col2", "Row1Col3"],
                ["Row2Col1", "Row2Col2", "Row2Col3"],
            ]
        )
        fake_docx = FakeDOCX()
        fake_docx.tables = [fake_table]
        loader = FakeDOCXLoader()
        loader.fake_docx = fake_docx
        extractor = DataExtractor(loader)
        data = extractor.extract_tables("dummy.docx")
        self.record_detail(
            f"Expected header: ['Header1', 'Header2', 'Header3'], got: {data[0]['data'][0]}"
        )
        self.assertEqual(data[0]["data"][0], ["Header1", "Header2", "Header3"])

    def test_ppt_extraction_text(self):
        extractor = DataExtractor(FakePPTLoader())
        data = extractor.extract_text("dummy.pptx")
        self.assertEqual(len(data), 2)
        slide1 = data[0]
        self.assertEqual(slide1["slide"], 1)
        self.assertTrue(any("text" in shape for shape in slide1["shapes"]))

    def test_ppt_extraction_tables(self):
        extractor = DataExtractor(FakePPTLoader())
        tables = extractor.extract_tables("dummy.pptx")
        self.assertGreaterEqual(len(tables), 1)
        self.assertEqual(tables[0]["data"][0], ["Name", "Grade", "Remark"])

    def test_ppt_extraction_images(self):
        extractor = DataExtractor(FakePPTLoaderWithImage())
        images = extractor.extract_images("dummy.pptx")
        self.record_detail(f"Extracted PPT image count: {len(images)}")
        self.assertGreaterEqual(len(images), 1)
        self.assertEqual(images[0]["format"], "jpeg")

    def test_invalid_file_loader(self):
        extractor = DataExtractor(InvalidFileLoader())
        with self.assertRaises(ValueError):
            extractor.extract_text("dummy.invalid")

    def test_corrupted_file(self):
        extractor = DataExtractor(CorruptedPDFLoader())
        with self.assertRaises(IOError):
            extractor.extract_text("corrupted.pdf")

    def test_storage_file(self):
        temp_dir = tempfile.TemporaryDirectory()
        storage = FileStorage(output_dir=temp_dir.name)
        sample_text = [{"page": 1, "line_text": "Test text"}]
        storage.store_data(sample_text, "text")
        file_path = os.path.join(temp_dir.name, "extracted_text.json")
        self.assertTrue(os.path.exists(file_path))
        temp_dir.cleanup()

    def test_storage_sql(self):
        db_fd, db_path = tempfile.mkstemp(suffix=".db")
        storage = SQLStorage(db_path=db_path)
        sample_text = [{"page": 1, "line_text": "Test text"}]
        storage.store_data(sample_text, "text")
        conn = sqlite3.connect(db_path)
        cursor = conn.cursor()
        cursor.execute("SELECT content FROM text_data")
        rows = cursor.fetchall()
        self.assertGreater(len(rows), 0)
        conn.close()
        os.close(db_fd)
        os.remove(db_path)


if __name__ == "__main__":
    runner = ExcelTestRunner(verbosity=2)
    unittest.main(testRunner=runner)
